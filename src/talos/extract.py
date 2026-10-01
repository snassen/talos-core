"""Text and facts out of attachments: PDF text, spreadsheet cells, photo EXIF.

Each extractor is bounded (pages, cells, characters), so one enormous file cannot
stall an ingest run. Anything unexpected is recorded as an error on the attachment
rather than raised: the original is safe in the vault and can be retried.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

MAX_CHARS = 200_000
MAX_PDF_PAGES = 60
MAX_CELLS = 20_000


@dataclass
class Extraction:
    status: str  # ok | empty | unsupported | encrypted | error
    text: str | None = None
    attrs: dict = field(default_factory=dict)


def _guess(content_type: str, filename: str | None) -> str:
    ct = (content_type or "").lower()
    name = (filename or "").lower()
    if ct == "application/pdf" or name.endswith(".pdf"):
        return "pdf"
    if "spreadsheetml" in ct or name.endswith((".xlsx", ".xlsm")):
        return "xlsx"
    if ct.startswith("image/") or name.endswith((".jpg", ".jpeg", ".png", ".heic", ".gif", ".webp", ".tif", ".tiff")):
        return "image"
    if ct in ("text/plain", "text/csv", "text/calendar") or name.endswith((".txt", ".csv", ".ics", ".log")):
        return "text"
    return "other"


_SURROGATES = re.compile("[\ud800-\udfff]")


def _clean(text: str) -> str:
    """PostgreSQL text cannot hold NUL, and psycopg cannot encode a lone surrogate."""
    return _SURROGATES.sub("\ufffd", text.replace("\x00", ""))


def extract(data: bytes, content_type: str, filename: str | None) -> Extraction:
    ex = _extract(data, content_type, filename)
    if ex.text:
        ex.text = _clean(ex.text)
    ex.attrs = {k: (_clean(v) if isinstance(v, str) else v) for k, v in ex.attrs.items()}
    return ex


def _extract(data: bytes, content_type: str, filename: str | None) -> Extraction:
    kind = _guess(content_type, filename)
    try:
        if kind == "pdf":
            return _pdf(data)
        if kind == "xlsx":
            return _xlsx(data)
        if kind == "image":
            return _image(data)
        if kind == "text":
            text = data.decode("utf-8", errors="replace")[:MAX_CHARS]
            return Extraction("ok" if text.strip() else "empty", text)
        return Extraction("unsupported")
    except Exception as exc:
        return Extraction("error", None, {"error": f"{type(exc).__name__}: {exc}"[:300]})


def _pdf(data: bytes) -> Extraction:
    from pypdf import PdfReader

    from pypdf import PasswordType

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        # decrypt("") opens a PDF that only has an owner password. It reports failure by its
        # return value, not by raising; a PDF that needs a password is marked 'encrypted'.
        try:
            opened = reader.decrypt("")
        except Exception as exc:
            return Extraction("encrypted", None, {"encrypted": True, "error": f"{type(exc).__name__}: {exc}"[:300]})
        if opened == PasswordType.NOT_DECRYPTED:
            return Extraction("encrypted", None, {"encrypted": True})
    pages = len(reader.pages)
    chunks, total = [], 0
    for page in reader.pages[:MAX_PDF_PAGES]:
        t = page.extract_text() or ""
        chunks.append(t)
        total += len(t)
        if total >= MAX_CHARS:
            break
    text = "\n".join(chunks)[:MAX_CHARS]
    attrs = {"pages": pages}
    meta = reader.metadata or {}
    for k in ("/Title", "/Author", "/Producer"):
        if meta.get(k):
            attrs[k.strip("/").lower()] = str(meta.get(k))[:200]
    return Extraction("ok" if text.strip() else "empty", text or None, attrs)


def _xlsx(data: bytes) -> Extraction:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines, cells = [], 0
    for ws in wb.worksheets:
        lines.append(f"# {ws.title}")
        for row in ws.iter_rows(values_only=True):
            values = [str(v) for v in row if v is not None and str(v).strip()]
            if values:
                lines.append("\t".join(values))
                cells += len(values)
            if cells >= MAX_CELLS:
                break
        if cells >= MAX_CELLS:
            break
    text = "\n".join(lines)[:MAX_CHARS]
    attrs = {"sheets": wb.sheetnames}
    wb.close()
    return Extraction("ok" if cells else "empty", text, attrs)


def _image(data: bytes) -> Extraction:
    from PIL import ExifTags, Image

    with Image.open(io.BytesIO(data)) as im:
        attrs = {"width": im.width, "height": im.height, "format": im.format}
        exif = im.getexif()
        if exif:
            named = {ExifTags.TAGS.get(k, str(k)): v for k, v in exif.items()}
            for key in ("DateTime", "Make", "Model", "Software"):
                if named.get(key):
                    attrs[key.lower()] = str(named[key])[:100]
            try:
                sub = exif.get_ifd(ExifTags.IFD.Exif)
                if sub.get(ExifTags.Base.DateTimeOriginal):
                    attrs["taken"] = str(sub[ExifTags.Base.DateTimeOriginal])
            except Exception:
                pass
    return Extraction("ok", None, attrs)
